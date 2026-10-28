# Registro de las cartas EMBLEMATICAS del juego (las famosas del anime / meta clasico, elegidas con el usuario
# 2026-10-16) y Excel con las 771 cartas usables marcadas Si/No. Las emblematicas se producen con cuidado una a una;
# las NO emblematicas de monstruo van en serie (TRELLIS por lotes).
# Salidas: remake/diseno/emblematicas.json y cartas-emblematicas.xlsx (raiz del proyecto, no se versiona).
# Uso: python -I herramientas/emblematicas.py   (openpyxl en D:/tools/pylib)
import csv, json, os, re, sys

sys.path.insert(0, "D:/tools/pylib")
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

EMBLEMATICAS = [
    # monstruos
    "Blue-Eyes White Dragon", "Red-Eyes B. Dragon", "Dark Magician", "Dark Magician Girl", "Summoned Skull",
    "Kuriboh", "Baby Dragon", "Celtic Guardian", "Silver Fang", "Exodia the Forbidden One",
    "Left Arm of the Forbidden One", "Right Arm of the Forbidden One", "Left Leg of the Forbidden One",
    "Right Leg of the Forbidden One", "Jinzo", "Buster Blader", "Gaia The Fierce Knight", "Time Wizard", "Great Moth",
    "Perfectly Ultimate Great Moth", "Insect Queen", "Harpie Lady", "Harpie Lady Sisters", "Harpie's Pet Dragon",
    "Gearfried the Iron Knight", "Panther Warrior", "Rocket Warrior", "Alligator's Sword", "Swordsman of Landstar",
    "Curse of Dragon", "Gemini Elf", "Vorse Raider", "Dark Elf", "Big Shield Gardna", "Labyrinth Wall", "Cyber Jar",
    "Sangan", "Witch of the Black Forest", "Magician of Faith", "Red-Eyes Black Metal Dragon",
    # fusiones
    "Blue-Eyes Ultimate Dragon", "Gaia the Dragon Champion", "Twin-Headed Thunder Dragon", "Thousand Dragon",
    "Flame Swordsman", "Alligator's Sword Dragon",
    # magias
    "Raigeki", "Monster Reborn", "Pot of Greed", "Dark Hole", "Swords of Revealing Light", "Polymerization",
    "Change of Heart", "Graceful Charity", "Heavy Storm", "Mystical Space Typhoon", "Harpie's Feather Duster",
    "Snatch Steal", "Premature Burial", "Megamorph", "Axe of Despair", "Black Pendant", "Fissure",
    "Nobleman of Crossout", "Scapegoat",
    # trampas
    "Mirror Force", "Magic Cylinder", "Trap Hole", "Call of the Haunted", "Torrential Tribute", "Ring of Destruction",
    "Solemn Judgment", "Magic Jammer", "Seven Tools of the Bandit", "Royal Decree", "Imperial Order", "Gravity Bind",
    "Dragon Capture Jar", "Waboku", "Negate Attack", "Horn of Heaven", "Dust Tornado", "Reinforcements",
]
# famosas que NO trae este juego (para no buscarlas)
NO_EN_EL_JUEGO = ["Black Skull Dragon", "Kaiser Sea Horse", "Dark Blade", "Fiber Jar", "Sakuretsu Armor",
                  "Bottomless Trap Hole", "Lightning Vortex", "Monster Reincarnation", "Little-Winguard"]
# carpeta de nano-banana/modelos3d cuyo nombre no coincide con el slug de la carta
CARPETA_MODELO = {"blue_eyes": "Blue-Eyes White Dragon", "red_eyes": "Red-Eyes B. Dragon",
                  "gaia": "Gaia The Fierce Knight"}


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def main():
    pack = open(os.path.join(BASE, "datadat-extraido", "bin#", "card_pack.bin"), "rb").read()
    usable = lambda r: int.from_bytes(pack[2 * int(r["idx"]):2 * int(r["idx"]) + 2], "little") & 7
    plantillas = json.load(open(os.path.join(BASE, "remake", "diseno", "plantillas.json"), encoding="utf-8"))["cartas"]
    modelos = {CARPETA_MODELO.get(d, d) for d in os.listdir(os.path.join(BASE, "nano-banana", "modelos3d"))}
    modelos = {slug(m) for m in modelos}
    figuras = {slug(os.path.splitext(f)[0]) for f in os.listdir(os.path.join(BASE, "nano-banana", "todo", "figura"))}
    emb = {e.lower() for e in EMBLEMATICAS}
    filas, vistos = [], set()
    for r in csv.DictReader(open(os.path.join(BASE, "re-output", "cards.csv"), encoding="utf-8")):
        if not r["idx"].isdigit() or not usable(r) or r["nombre"].lower() in vistos:
            continue
        vistos.add(r["nombre"].lower())
        s = slug(r["nombre"])
        clase = {"Magic": "Magia", "Trap": "Trampa"}.get(r["tipo"], "Monstruo")
        p = plantillas.get(s, {})
        filas.append({"carta": r["nombre"], "clase": clase, "emblematica": r["nombre"].lower() in emb,
                      "tipo": "" if clase != "Monstruo" else r["tipo"], "atributo": "" if clase != "Monstruo" else r["attr"],
                      "nivel": "" if clase != "Monstruo" else int(r["nivel"] or 0),
                      "atk": "" if clase != "Monstruo" else int(r["atk"] or 0),
                      "def": "" if clase != "Monstruo" else int(r["defn"] or 0),
                      "plantilla": p.get("plantilla", ""), "modelo_3d": s in modelos,
                      "figura_gemini": clase == "Monstruo" and s in figuras})
    faltan = sorted(emb - {f["carta"].lower() for f in filas})
    if faltan:
        print("AVISO emblematicas que no estan entre las usables:", faltan)
    filas.sort(key=lambda f: (not f["emblematica"], ["Monstruo", "Magia", "Trampa"].index(f["clase"]), f["carta"]))
    json.dump({"_nota": "Cartas emblematicas (elegidas con el usuario 2026-10-16). Las NO emblematicas de monstruo se "
               "producen en serie. Generado por herramientas/emblematicas.py",
               "emblematicas": [f["carta"] for f in filas if f["emblematica"]],
               "no_emblematicas_monstruo": [f["carta"] for f in filas if not f["emblematica"] and f["clase"] == "Monstruo"],
               "no_estan_en_el_juego": NO_EN_EL_JUEGO},
              open(os.path.join(BASE, "remake", "diseno", "emblematicas.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cartas"
    cab = ["Carta", "Clase", "Emblemática", "Tipo", "Atributo", "Nivel", "ATK", "DEF", "Plantilla 3D", "Modelo 3D hecho", "Figura Gemini"]
    ws.append(cab)
    oro = PatternFill("solid", fgColor="FFE699")
    for f in filas:
        ws.append([f["carta"], f["clase"], "Sí" if f["emblematica"] else "No", f["tipo"], f["atributo"], f["nivel"],
                   f["atk"], f["def"], f["plantilla"], "Sí" if f["modelo_3d"] else "No",
                   "" if f["clase"] != "Monstruo" else ("Sí" if f["figura_gemini"] else "No")])
        if f["emblematica"]:
            for c in ws[ws.max_row]:
                c.fill = oro
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center")
    for col, ancho in zip("ABCDEFGHIJK", (42, 11, 13, 14, 10, 7, 7, 7, 20, 16, 14)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    r = wb.create_sheet("Resumen")
    r.append(["", "Emblemáticas", "No emblemáticas", "Total"])
    for clase in ("Monstruo", "Magia", "Trampa"):
        e = sum(1 for f in filas if f["clase"] == clase and f["emblematica"])
        t = sum(1 for f in filas if f["clase"] == clase)
        r.append([clase, e, t - e, t])
    e = sum(1 for f in filas if f["emblematica"])
    r.append(["Total", e, len(filas) - e, len(filas)])
    r.append([])
    r.append(["Famosas que NO trae este juego:", ", ".join(NO_EN_EL_JUEGO)])
    for c in r[1]:
        c.font = Font(bold=True)
    r.column_dimensions["A"].width = 32
    for col in "BCD":
        r.column_dimensions[col].width = 18
    salida = os.path.join(BASE, "cartas-emblematicas.xlsx")
    wb.save(salida)
    for fila in r.iter_rows(min_row=2, max_row=5, values_only=True):
        print(fila)
    print("guardado", salida)


main()
